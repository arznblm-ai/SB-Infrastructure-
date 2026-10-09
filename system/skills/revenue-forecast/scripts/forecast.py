#!/usr/bin/env python3
"""Bottom-up помесячный прогноз выручки B2B-подписки (пилоты -> платящие -> MRR).

Механика на месяц m (1..months), по каждому сценарию:
  pilots_m      = pilots_by_halfyear[min(len-1,(m-1)//6)]
                  + seller_pilots * min(1, (m-hire_month+1)/ramp_months)   (если m >= hire_month)
                  -> min(pilots_m, impl_capacity)
  new_paying_m  = pilots_{m-pilot_months} * conversion
  clients_m     = clients_{m-1} * (1 - churn_monthly) + new_paying_m
  check         = license_base + max(0, users-included_users)*extra_user_price + support_monthly
                  (или users*price_per_seat + support_monthly, если задан price_per_seat)
  mrr_m         = clients_m * check
  hours_m       = new_paying_m*onboarding_hours + clients_m*support_hours_monthly + pilots_m*pilot_hours
  cost          = hours * hour_cost ; profit = revenue - cost - revenue*tax

Параметры сценария перекрывают "defaults" и "pricing" из JSON.

Примеры:
  python3 forecast.py example-zekiro.json
  python3 forecast.py example-zekiro.json --csv /tmp/out.csv --xlsx /tmp/out.xlsx
  python3 forecast.py example-zekiro.json --sensitivity Базовый
"""
import argparse
import csv
import json
import sys

REQUIRED = ["pilots_by_halfyear", "conversion", "pilot_months", "churn_monthly",
            "users", "impl_capacity"]
DEFAULTS = {"ramp_months": 6, "support_monthly": 0, "hire_month": None,
            "seller_pilots": 0, "onboarding_hours": 0, "support_hours_monthly": 0,
            "pilot_hours": 0, "team_hours_capacity": None, "hour_cost": None, "tax": 0,
            "license_base": 0, "included_users": 0, "extra_user_price": 0,
            "price_per_seat": None}


class ConfigError(Exception):
    pass


# ---------- форматирование ----------

def fnum(x, digits=0):
    if x is None:
        return "—"
    s = f"{x:,.{digits}f}".replace(",", " ")
    return s.replace(".", ",")


def fmoney(x, cur):
    return f"{fnum(x)} {cur}"


# ---------- конфиг ----------

def load_config(path):
    try:
        with open(path, encoding="utf-8") as f:
            cfg = json.load(f)
    except FileNotFoundError:
        raise ConfigError(f"Файл не найден: {path}")
    except json.JSONDecodeError as e:
        raise ConfigError(f"Невалидный JSON в {path}: {e}")
    if not isinstance(cfg.get("scenarios"), dict) or not cfg["scenarios"]:
        raise ConfigError("В JSON нет непустого объекта 'scenarios'")
    return cfg


def resolve(cfg, name):
    if name not in cfg["scenarios"]:
        raise ConfigError(f"Сценарий '{name}' не найден. Есть: {', '.join(cfg['scenarios'])}")
    p = dict(DEFAULTS)
    p.update(cfg.get("pricing", {}))
    p.update(cfg.get("defaults", {}))
    p.update(cfg["scenarios"][name])
    missing = [k for k in REQUIRED if p.get(k) is None]
    if p.get("price_per_seat") is None and not p.get("license_base"):
        missing.append("license_base (или price_per_seat)")
    if missing:
        raise ConfigError(f"Сценарий '{name}': не хватает полей: {', '.join(missing)}")
    if not isinstance(p["pilots_by_halfyear"], list) or not p["pilots_by_halfyear"]:
        raise ConfigError(f"Сценарий '{name}': pilots_by_halfyear должен быть непустым списком")
    if p["hire_month"] is not None and p["ramp_months"] <= 0:
        raise ConfigError(f"Сценарий '{name}': ramp_months должен быть > 0")
    return p


# ---------- модель ----------

def check_price(p):
    if p.get("price_per_seat") is not None:
        lic = p["users"] * p["price_per_seat"]
    else:
        lic = p["license_base"] + max(0, p["users"] - p["included_users"]) * p["extra_user_price"]
    return lic + p["support_monthly"]


def simulate(p, months):
    pm = int(round(p["pilot_months"]))
    hb = p["pilots_by_halfyear"]
    chk = check_price(p)
    rows, pilots_hist = [], []
    clients = rev_cum = 0.0
    for m in range(1, months + 1):
        pilots = hb[min(len(hb) - 1, (m - 1) // 6)]
        hm = p["hire_month"]
        if hm is not None and m >= hm:
            pilots += p["seller_pilots"] * min(1.0, (m - hm + 1) / p["ramp_months"])
        pilots = min(pilots, p["impl_capacity"])
        pilots_hist.append(pilots)
        src = m - pm
        new_paying = pilots_hist[src - 1] * p["conversion"] if src >= 1 else 0.0
        clients = clients * (1 - p["churn_monthly"]) + new_paying
        mrr = clients * chk
        rev_cum += mrr
        hours = (new_paying * p["onboarding_hours"] + clients * p["support_hours_monthly"]
                 + pilots * p["pilot_hours"])
        rows.append(dict(month=m, pilots=pilots, new_paying=new_paying, clients=clients,
                         mrr=mrr, revenue_cum=rev_cum, hours=hours))
    return rows


def summarize(p, rows, fx):
    n = len(rows)

    def at(m, k):
        return rows[min(m, n) - 1][k]

    y1 = sum(r["mrr"] for r in rows[:12])
    y2 = sum(r["mrr"] for r in rows[12:24])
    total = rows[-1]["revenue_cum"]
    hours_total = sum(r["hours"] for r in rows)
    peak_h = max(r["hours"] for r in rows)
    cap = p.get("team_hours_capacity")
    peak_load = peak_h / cap * 100 if cap else None
    over = [r["month"] for r in rows if cap and r["hours"] > cap]
    profit = None
    if p.get("hour_cost") is not None:
        profit = total - hours_total * p["hour_cost"] - total * p.get("tax", 0)
    arr = rows[-1]["mrr"] * 12
    return dict(check=check_price(p), clients12=at(12, "clients"), clients24=at(24, "clients"),
                mrr12=at(12, "mrr"), mrr24=at(24, "mrr"), rev_y1=y1, rev_y2=y2, total=total,
                arr=arr, arr_usd=arr / fx if fx else None,
                pilots=sum(r["pilots"] for r in rows),
                paying=sum(r["new_paying"] for r in rows),
                peak_hours=peak_h, peak_load=peak_load, over_months=over, profit=profit)


# ---------- вывод ----------

SUMMARY_LINES = [
    ("Чек клиента / мес", "check", "money"),
    ("Платящих к 12 мес", "clients12", "num1"),
    ("Платящих к 24 мес", "clients24", "num1"),
    ("MRR на 12 мес", "mrr12", "money"),
    ("MRR на 24 мес", "mrr24", "money"),
    ("Выручка 1-й год", "rev_y1", "money"),
    ("Выручка 2-й год", "rev_y2", "money"),
    ("Выручка за горизонт", "total", "money"),
    ("ARR на конец горизонта", "arr", "money"),
    ("ARR на конец, $", "arr_usd", "usd"),
    ("Пилотов за горизонт", "pilots", "num1"),
    ("Пришло платящих", "paying", "num1"),
    ("Пик часов команды / мес", "peak_hours", "num0"),
    ("Пик загрузки команды", "peak_load", "pct"),
    ("Месяцы перегруза >100%", "over_months", "list"),
    ("Прибыль за горизонт", "profit", "money"),
]


def fmt_cell(v, kind, cur):
    if v is None:
        return "—"
    if kind == "money":
        return fmoney(v, cur)
    if kind == "usd":
        return f"$ {fnum(v)}"
    if kind == "num1":
        return fnum(v, 1)
    if kind == "num0":
        return fnum(v)
    if kind == "pct":
        return f"{fnum(v)} %"
    if kind == "list":
        return ", ".join(map(str, v)) if v else "нет"
    return str(v)


def print_table(header, body):
    widths = [max(len(str(r[i])) for r in [header] + body) for i in range(len(header))]
    line = lambda r: "  ".join(str(c).ljust(widths[0]) if i == 0 else str(c).rjust(widths[i])
                               for i, c in enumerate(r))
    print(line(header))
    print("  ".join("-" * w for w in widths))
    for r in body:
        print(line(r))


def print_summary(names, sums, cur, months):
    print(f"Прогноз выручки, горизонт {months} мес, валюта {cur}\n")
    body = [[label] + [fmt_cell(sums[n][k], kind, cur) for n in names]
            for label, k, kind in SUMMARY_LINES]
    print_table(["Показатель"] + names, body)
    for n in names:
        if sums[n]["over_months"]:
            print(f"\n! {n}: загрузка команды > 100% в месяцах {fmt_cell(sums[n]['over_months'], 'list', cur)}")


def write_csv(path, all_rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["scenario", "month", "pilots", "new_paying", "clients", "mrr",
                    "revenue_cum", "hours"])
        for name, rows in all_rows.items():
            for r in rows:
                w.writerow([name, r["month"], round(r["pilots"], 4), round(r["new_paying"], 4),
                            round(r["clients"], 4), round(r["mrr"], 2),
                            round(r["revenue_cum"], 2), round(r["hours"], 2)])


def write_xlsx(path, names, sums, all_rows, cur):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise ConfigError("Для --xlsx нужен openpyxl: pip install openpyxl")
    reg, bold = Font(name="Arial"), Font(name="Arial", bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Показатель"] + names)
    for label, k, kind in SUMMARY_LINES:
        vals = []
        for n in names:
            v = sums[n][k]
            if kind == "list":
                v = ", ".join(map(str, v)) if v else "нет"
            elif isinstance(v, float):
                v = round(v, 2)
            vals.append(v)
        ws.append([label] + vals)
    ws.append([])
    ws.append([f"Валюта: {cur}; пик загрузки — в % от team_hours_capacity"])
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 18
    for name in names:
        sh = wb.create_sheet(title=name[:31])
        sh.append(["Месяц", "Новые пилоты", "Новые платящие", "Клиенты", f"MRR, {cur}",
                   f"Выручка накоп., {cur}", "Часы команды"])
        for r in all_rows[name]:
            sh.append([r["month"], round(r["pilots"], 3), round(r["new_paying"], 3),
                       round(r["clients"], 3), round(r["mrr"], 2), round(r["revenue_cum"], 2),
                       round(r["hours"], 2)])
        for col in "ABCDEFG":
            sh.column_dimensions[col].width = 18
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                c.font = bold if c.row == 1 else reg
    wb.save(path)


# ---------- чувствительность ----------

def sensitivity(cfg, name, months):
    base = resolve(cfg, name)
    base_total = simulate(base, months)[-1]["revenue_cum"]

    def run(over):
        p = dict(base)
        p.update(over)
        return simulate(p, months)[-1]["revenue_cum"]

    hb = base["pilots_by_halfyear"]
    tests = [
        ("conversion ×0,7/×1,3", {"conversion": base["conversion"] * 0.7},
         {"conversion": min(1.0, base["conversion"] * 1.3)}),
        ("pilots_by_halfyear ×0,7/×1,3", {"pilots_by_halfyear": [x * 0.7 for x in hb]},
         {"pilots_by_halfyear": [x * 1.3 for x in hb]}),
        ("users ±30%", {"users": base["users"] * 0.7}, {"users": base["users"] * 1.3}),
        ("churn_monthly ×0,5/×2", {"churn_monthly": base["churn_monthly"] * 0.5},
         {"churn_monthly": base["churn_monthly"] * 2}),
        ("pilot_months ±1", {"pilot_months": max(0, base["pilot_months"] - 1)},
         {"pilot_months": base["pilot_months"] + 1}),
    ]
    if base["hire_month"] is not None:
        tests.append(("hire_month ±3 мес", {"hire_month": max(1, base["hire_month"] - 3)},
                      {"hire_month": base["hire_month"] + 3}))
    res = []
    for label, lo, hi in tests:
        a, b = run(lo), run(hi)
        res.append((label, a, b, abs(b - a)))
    res.sort(key=lambda x: -x[3])
    cur = cfg.get("currency", "₽")
    print(f"Tornado: «{name}», выручка за {months} мес. База: {fmoney(base_total, cur)}\n")
    body = [[label, fmoney(a, cur), fmoney(b, cur), fmoney(sp, cur),
             f"{fnum((min(a, b) - base_total) / base_total * 100)}…+{fnum((max(a, b) - base_total) / base_total * 100)} %"]
            for label, a, b, sp in res]
    print_table(["Драйвер", "Нижнее", "Верхнее", "Размах", "К базе"], body)
    print("\n(«Нижнее»/«Верхнее» — выручка при первом/втором значении драйвера в подписи)")


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser(
        description="Bottom-up помесячный прогноз выручки (пилоты → платящие → MRR) по сценариям из JSON.",
        epilog="Пример: python3 forecast.py example-zekiro.json --sensitivity Базовый")
    ap.add_argument("config", help="JSON с вводными (см. example-zekiro.json)")
    ap.add_argument("--csv", metavar="OUT.csv", help="помесячная таблица всех сценариев в CSV")
    ap.add_argument("--xlsx", metavar="OUT.xlsx", help="книга Excel: Сводка + лист на сценарий")
    ap.add_argument("--sensitivity", metavar="SCENARIO", help="tornado по драйверам для сценария")
    ap.add_argument("--months", type=int, help="горизонт, мес (перекрывает months из JSON)")
    a = ap.parse_args()
    try:
        cfg = load_config(a.config)
        months = a.months or cfg.get("months", 24)
        if months < 1:
            raise ConfigError("months должен быть >= 1")
        cur, fx = cfg.get("currency", "₽"), cfg.get("fx")
        if a.sensitivity:
            sensitivity(cfg, a.sensitivity, months)
            return 0
        names = list(cfg["scenarios"])
        params = {n: resolve(cfg, n) for n in names}
        all_rows = {n: simulate(params[n], months) for n in names}
        sums = {n: summarize(params[n], all_rows[n], fx) for n in names}
        print_summary(names, sums, cur, months)
        if a.csv:
            write_csv(a.csv, all_rows)
            print(f"\nCSV: {a.csv}")
        if a.xlsx:
            write_xlsx(a.xlsx, names, sums, all_rows, cur)
            print(f"XLSX: {a.xlsx}")
    except ConfigError as e:
        print(f"Ошибка: {e}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
