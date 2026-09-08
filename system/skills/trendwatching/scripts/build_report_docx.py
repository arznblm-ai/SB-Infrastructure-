#!/usr/bin/env python3
"""Собирает Word-отчёт по трендам из trends.json.

    python3 build_report_docx.py trends.json -o "Тренд-отчёт.docx"
    python3 build_report_docx.py trends.json -o report.docx --lang en

Язык берётся из --lang, иначе из meta.lang, иначе русский.
Структура разделов зафиксирована в references/07-delivery.md.
Требует python-docx: pip install python-docx --break-system-packages
"""

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import trendlib as T  # noqa: E402

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
except ImportError:
    sys.exit("Нужен python-docx: pip install python-docx --break-system-packages")

ACCENT = RGBColor(0x1E, 0x3F, 0xCB)
MUTED = RGBColor(0x66, 0x71, 0x6A)

S = {
    "kicker": ("ТРЕНД-ОТЧЁТ", "TREND REPORT"),
    "untitled": ("Тренды", "Trends"),
    "field": ("Поле", "Field"),
    "value": ("Значение", "Value"),
    "horizon": ("Горизонт", "Horizon"),
    "geo": ("География", "Geography"),
    "segments": ("Сегменты", "Segments"),
    "depth": ("Глубина прогона", "Scan depth"),
    "wave": ("Волна", "Wave"),
    "date": ("Дата", "Date"),
    "author": ("Автор", "Author"),
    "profile": ("Профиль весов", "Weight profile"),
    "howto": ("Как читать этот отчёт", "How to read this report"),
    "howto_fia": (
        "[F] — наблюдаемый факт: проверяемое утверждение с датой, источником и методикой. "
        "[I] — интерпретация: что это может означать. [A] — допущение: непроверенная предпосылка, "
        "необходимая для решения. Смешение этих уровней — самая частая методологическая ошибка, "
        "поэтому они разделены явно.",
        "[F] is an observed fact: a verifiable statement with a date, a source and a method. "
        "[I] is interpretation: what it may mean. [A] is an assumption: an unverified premise a decision "
        "rests on. Conflating these layers is the most common methodological error, so they are separated "
        "explicitly."),
    "howto_rel": (
        "Классы надёжности источников: A — официальные, peer-reviewed, нормативные; "
        "B — первичные корпоративные или авторитетные профессиональные; C — качественный вторичный контекст.",
        "Source reliability classes: A — official, peer-reviewed, statutory; B — primary corporate or "
        "authoritative professional; C — qualitative secondary context."),
    "howto_score": (
        "Score — оценка доказательности по 15 взвешенным критериям (0–100). Coverage — доля весов, "
        "по которым есть фактическая оценка. Неизвестные критерии помечены N/D и не заменяются нулём: "
        "нехватка данных снижает уверенность, а не балл.",
        "Score is an evidence rating across 15 weighted criteria (0–100). Coverage is the share of weight "
        "actually assessed. Unknown criteria are marked N/D and never silently set to zero: missing data "
        "lowers confidence, not the score."),
    "limits_default": (
        "Отчёт не является прогнозом, юридическим заключением или оценкой размера рынка. "
        "Актуальные решения требуют свежих локальных данных.",
        "This report is not a forecast, a legal opinion or a market sizing. Live decisions require fresh "
        "local data."),
    "frame": ("Рамка: какие решения этот отчёт должен улучшить", "Frame: which decisions this report should improve"),
    "question": ("Вопрос", "Question"),
    "owner": ("Владелец", "Owner"),
    "decision_date": ("Дата решения", "Decision date"),
    "boundaries": ("Границы", "Boundaries"),
    "summary": ("Резюме", "Executive summary"),
    "map": ("Карта трендов", "Trend map"),
    "trend": ("Тренд", "Trend"),
    "stage": ("Стадия", "Stage"),
    "score": ("Score", "Score"),
    "cov": ("Cov.", "Cov."),
    "conf": ("Увер.", "Conf."),
    "pi": ("P × I", "P × I"),
    "decision": ("Решение", "Decision"),
    "cards": ("Карточки трендов", "Trend cards"),
    "metric": ("Показатель", "Metric"),
    "score_cov_conf": ("Score / coverage / уверенность", "Score / coverage / confidence"),
    "status": ("Статус", "Status"),
    "prob_impact": ("Вероятность × влияние", "Probability × impact"),
    "object": ("объект", "object"),
    "urgency_window": ("Срочность / окно решения", "Urgency / decision window"),
    "domains": ("Домены", "Domains"),
    "hgs": ("Горизонт / география / сегменты", "Horizon / geography / segments"),
    "independent": ("Независимых сигналов", "Independent signals"),
    "owner_review": ("Владелец / следующий пересмотр", "Owner / next review"),
    "prob_cond": ("Условие оценки вероятности", "Probability is conditional on"),
    "excluded": ("Что не входит", "Out of scope"),
    "drivers": ("Драйверы", "Drivers"),
    "enablers": ("Enablers", "Enablers"),
    "barriers": ("Барьеры", "Barriers"),
    "needs": ("Базовые потребности", "Basic needs"),
    "shifts": ("Долгосрочные сдвиги", "Long-term shifts"),
    "triggers": ("Краткосрочные триггеры", "Short-term triggers"),
    "gap": ("Разрыв ожиданий", "Expectation gap"),
    "curves": ("Пять кривых (0–5)", "Five curves (0–5)"),
    "counter": ("Контртренд", "Counter-trend"),
    "disconf": ("Опровергающие свидетельства", "Disconfirming evidence"),
    "disconf_missing": (
        "Опровергающие свидетельства не найдены — это почти всегда признак недостаточного поиска, "
        "а не их отсутствия.",
        "No disconfirming evidence found — almost always a sign of insufficient search rather than its absence."),
    "cost": ("Цена принятия", "Cost of adoption"),
    "opps": ("Возможности", "Opportunities"),
    "risks": ("Риски", "Risks"),
    "hyps": ("Гипотезы", "Hypotheses"),
    "exp": ("Эксперимент и правило остановки", "Experiment and stop-rule"),
    "stop": ("Stop-rule", "Stop-rule"),
    "indicators": ("Индикаторы наблюдения", "Monitoring indicators"),
    "ind_name": ("Индикатор", "Indicator"),
    "ind_kind": ("Тип", "Type"),
    "ind_source": ("Источник", "Source"),
    "ind_base": ("База", "Baseline"),
    "ind_thr": ("Порог", "Threshold"),
    "leading": ("ведущий", "leading"),
    "lagging": ("отстающий", "lagging"),
    "ind_missing": (
        "Индикаторы не заданы — тренд нельзя ни подтвердить, ни опровергнуть.",
        "No indicators set — the trend can be neither confirmed nor refuted."),
    "gaps": ("Пробелы в оценке (N/D)", "Assessment gaps (N/D)"),
    "signals": ("Сигналы", "Signals"),
    "excl_signals": ("Рассмотрены и исключены", "Reviewed and excluded"),
    "delta": ("Изменение с прошлой волны", "Change since the previous wave"),
    "rejected": ("Что не подтвердилось", "What did not hold up"),
    "rejected_intro": (
        "Кандидаты, рассмотренные и отклонённые в этой волне. Раздел существует потому, "
        "что видимость отбраковки — главный признак того, что отбор вообще был.",
        "Candidates reviewed and rejected in this wave. The section exists because visible rejection is the "
        "main evidence that selection actually happened."),
    "rejected_none": (
        "Отклонённых кандидатов не зафиксировано. Если прогон был полным, это стоит перепроверить: "
        "отбор без отбраковки обычно означает, что кандидатов было слишком мало.",
        "No rejected candidates recorded. For a full scan this is worth re-checking: selection without "
        "rejection usually means there were too few candidates."),
    "candidate": ("Кандидат", "Candidate"),
    "reason": ("Основание", "Reason"),
    "appA": ("Приложение A. Реестр сигналов", "Appendix A. Signal register"),
    "appB": ("Приложение B. Scoring sheets", "Appendix B. Scoring sheets"),
    "appC": ("Приложение C. Источники и их ограничения", "Appendix C. Sources and their limits"),
    "obs": ("Наблюдение", "Observation"),
    "type": ("Тип", "Type"),
    "rel": ("Надёж.", "Rel."),
    "tier": ("Уровень", "Tier"),
    "link": ("Ссылка", "Link"),
    "review": ("Review", "Review"),
    "criterion": ("Критерий", "Criterion"),
    "weight": ("Вес", "Weight"),
    "rating": ("Оценка", "Rating"),
    "contrib": ("Вклад", "Contribution"),
    "fact_ref": ("Факт или ссылка", "Fact or reference"),
    "normalized": ("нормированный", "normalised"),
    "src_type": ("Тип источника", "Source type"),
    "src_used": ("Использованные источники", "Sources used"),
    "src_count": ("Сигналов", "Signals"),
    "appC_note": (
        "Проверьте перед публикацией: не доминирует ли один тип источника (порог ~35%), "
        "какие аудитории и регионы отсутствуют, и не считаются ли перепечатки независимыми подтверждениями.",
        "Check before publishing: whether one source type dominates (threshold ~35%), which audiences and "
        "regions are missing, and whether reprints are being counted as independent confirmations."),
    "tier_note": ("Уровни доказательности", "Evidence tiers"),
}


def t(key, lang):
    return T.L(S[key], lang)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def para(doc, text, *, size=10.5, bold=False, italic=False, color=None,
         space_before=0, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def label(doc, text):
    para(doc, text.upper(), size=8, bold=True, color=MUTED, space_before=10, space_after=2)


def table(doc, headers, rows, widths=None):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.style = "Light Grid Accent 1"
    for i, head in enumerate(headers):
        cell = tb.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(8.5)
    for row in rows:
        cells = tb.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if value is None else str(value))
            run.font.size = Pt(8.5)
    if widths:
        for row in tb.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return tb


def fmt_pi(trend):
    p, i = trend.get("probability"), trend.get("impact")
    if p is None or i is None:
        return "—"
    return f"{p} × {i} = {p * i}"


def build(data, out_path, lang="ru"):
    doc = Document()
    style_document(doc)
    meta = data["meta"]
    index = T.signals_index(data)
    profile, weight_table = T.weights(data)

    def L(pair):
        return T.L(pair, lang)

    # --- Титул -----------------------------------------------------------
    para(doc, t("kicker", lang), size=9, bold=True, color=MUTED, space_after=4)
    para(doc, meta.get("title", t("untitled", lang)), size=26, bold=True, space_after=10)
    para(doc, meta.get("object", ""), size=12, color=MUTED, space_after=14)
    rows = [
        [t("horizon", lang), meta.get("horizon", "—")],
        [t("geo", lang), meta.get("geo", "—")],
        [t("segments", lang), ", ".join(meta.get("segments", [])) or "—"],
        [t("depth", lang), L(T.DEPTHS.get(meta.get("depth"), ("—", "—")))],
        [t("wave", lang), meta.get("wave", 1)],
        [t("date", lang), meta.get("date", "—")],
        [t("author", lang), meta.get("author", "—")],
    ]
    if profile != "default":
        rows.append([t("profile", lang), profile])
    table(doc, [t("field", lang), t("value", lang)], rows, widths=[4.5, 11.5])

    # --- Как читать ------------------------------------------------------
    doc.add_heading(t("howto", lang), level=1)
    para(doc, t("howto_fia", lang))
    para(doc, t("howto_rel", lang))
    para(doc, t("howto_score", lang))
    para(doc, meta.get("limitations") or t("limits_default", lang), italic=True, color=MUTED)

    # --- Рамка -----------------------------------------------------------
    dqs = meta.get("decision_questions") or []
    if dqs:
        doc.add_heading(t("frame", lang), level=1)
        table(doc, [t("question", lang), t("owner", lang), t("decision_date", lang)],
              [[q.get("question", ""), q.get("owner", "—"), q.get("decision_date", "—")] for q in dqs],
              widths=[9.5, 3.5, 3.0])
    if meta.get("boundaries"):
        para(doc, f"{t('boundaries', lang)}: {meta['boundaries']}", color=MUTED)

    # --- Резюме ----------------------------------------------------------
    doc.add_heading(t("summary", lang), level=1)
    trends = sorted(data["trends"], key=lambda tr: (T.compute(tr, weight_table)["score"] or 0), reverse=True)
    for tr in trends[:6]:
        c = T.compute(tr, weight_table)
        status, _ = T.status_for(c["score"], lang)
        para(doc, f"{tr['name']} — {tr.get('statement', '')}", bold=True, space_after=1)
        para(doc, f"Score {c['score']} · coverage {c['coverage']}% · "
                  f"{L(T.CONFIDENCE[T.confidence_hint(tr, c, index)])} · {status} · "
                  f"{t('decision', lang).lower()}: {L(T.DECISIONS.get(tr.get('decision'), ('—', '—')))}",
             size=9, color=MUTED, space_after=8)

    # --- Карта трендов ---------------------------------------------------
    doc.add_heading(t("map", lang), level=1)
    rows = []
    for tr in trends:
        c = T.compute(tr, weight_table)
        rows.append([
            tr["name"], L(T.STAGES[tr.get("stage", 0)]), c["score"], f"{c['coverage']}%",
            L(T.CONFIDENCE[T.confidence_hint(tr, c, index)]), fmt_pi(tr),
            L(T.DECISIONS.get(tr.get("decision"), ("—", "—"))), tr.get("owner") or "—",
        ])
    table(doc, [t("trend", lang), t("stage", lang), t("score", lang), t("cov", lang),
                t("conf", lang), t("pi", lang), t("decision", lang), t("owner", lang)],
          rows, widths=[3.6, 2.6, 1.3, 1.2, 1.4, 1.6, 2.4, 1.9])

    # --- Карточки --------------------------------------------------------
    doc.add_page_break()
    doc.add_heading(t("cards", lang), level=1)
    for tr in trends:
        c = T.compute(tr, weight_table)
        conf = T.confidence_hint(tr, c, index)
        status, action = T.status_for(c["score"], lang)

        doc.add_heading(tr["name"], level=2)
        para(doc, tr.get("statement", ""), size=11.5, space_after=8)

        table(doc, [t("metric", lang), t("value", lang)], [
            [t("stage", lang), L(T.STAGES[tr.get("stage", 0)])],
            [t("score_cov_conf", lang), f"{c['score']} / {c['coverage']}% / {L(T.CONFIDENCE[conf])}"],
            [t("status", lang), f"{status} → {action}"],
            [t("prob_impact", lang), f"{fmt_pi(tr)} ({t('object', lang)}: {tr.get('impact_object') or '—'})"],
            [t("urgency_window", lang), f"{tr.get('urgency') or '—'} / {tr.get('decision_window') or '—'}"],
            [t("decision", lang), L(T.DECISIONS.get(tr.get("decision"), ("—", "—")))],
            [t("domains", lang), ", ".join(L(T.STEEP[s]) for s in tr.get("steep", []) if s in T.STEEP) or "—"],
            [t("hgs", lang), f"{tr.get('horizon') or '—'} · {tr.get('geo') or '—'} · "
                             f"{', '.join(tr.get('segments', [])) or '—'}"],
            [t("independent", lang), T.independent_count(tr, index)],
            [t("owner_review", lang), f"{tr.get('owner') or '—'} · {tr.get('next_review') or '—'}"],
        ], widths=[5.5, 10.5])

        if tr.get("probability_condition"):
            para(doc, f"{t('prob_cond', lang)}: {tr['probability_condition']}",
                 size=9, italic=True, color=MUTED)

        for key, title in [("boundaries", "boundaries"), ("excluded", "excluded")]:
            if tr.get(key):
                label(doc, t(title, lang))
                para(doc, tr[key])

        for key, title in [("drivers", "drivers"), ("enablers", "enablers"), ("barriers", "barriers"),
                           ("basic_needs", "needs"), ("shifts", "shifts"), ("triggers", "triggers")]:
            values = tr.get(key) or []
            if values:
                label(doc, t(title, lang))
                for v in values:
                    doc.add_paragraph(v, style="List Bullet")

        if tr.get("expectation_gap"):
            label(doc, t("gap", lang))
            para(doc, tr["expectation_gap"])

        curves = tr.get("curves") or {}
        if any(curves.get(k) is not None for k, _, _ in T.CURVES):
            label(doc, t("curves", lang))
            table(doc, [T.L((ru, en), lang) for _, ru, en in T.CURVES],
                  [[curves.get(k) if curves.get(k) is not None else "N/D" for k, _, _ in T.CURVES]])

        if tr.get("counter_trend"):
            label(doc, t("counter", lang))
            para(doc, tr["counter_trend"])
        label(doc, t("disconf", lang))
        if tr.get("disconfirming"):
            para(doc, tr["disconfirming"])
        else:
            para(doc, t("disconf_missing", lang), size=9, italic=True, color=MUTED)
        if tr.get("adoption_cost"):
            label(doc, t("cost", lang))
            para(doc, tr["adoption_cost"])

        for key, title in [("opportunities", "opps"), ("risks", "risks"), ("hypotheses", "hyps")]:
            values = tr.get(key) or []
            if values:
                label(doc, t(title, lang))
                for v in values:
                    doc.add_paragraph(v, style="List Bullet")

        if tr.get("experiment") or tr.get("stop_rule"):
            label(doc, t("exp", lang))
            para(doc, tr.get("experiment") or "—")
            para(doc, f"{t('stop', lang)}: {tr.get('stop_rule') or '—'}", size=9, color=MUTED)

        indicators = tr.get("indicators") or []
        if indicators:
            label(doc, t("indicators", lang))
            table(doc, [t("ind_name", lang), t("ind_kind", lang), t("ind_source", lang),
                        t("ind_base", lang), t("ind_thr", lang)],
                  [[i.get("name", ""),
                    t("leading", lang) if i.get("kind") == "leading" else t("lagging", lang),
                    i.get("source", "—"), i.get("baseline") or "—", i.get("threshold") or "—"]
                   for i in indicators], widths=[4.5, 2.2, 4.0, 2.6, 2.7])
        else:
            para(doc, t("ind_missing", lang), size=9, italic=True, color=MUTED)

        if c["gaps"]:
            para(doc, f"{t('gaps', lang)}: " + ", ".join(L(g) for g in c["gaps"]), size=9, color=MUTED)

        label(doc, t("signals", lang))
        for sid in tr.get("signal_ids", []):
            sig = index.get(sid)
            if not sig:
                continue
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{sid} · {sig.get('date_event', '')} · {T.tier(sig)} · {sig.get('headline_fact', '')}")
            r.font.size = Pt(9)
            if sig.get("url"):
                r2 = p.add_run(f"  {sig['url']}")
                r2.font.size = Pt(8)
                r2.font.color.rgb = ACCENT
        if tr.get("excluded_signal_ids"):
            para(doc, f"{t('excl_signals', lang)}: " + ", ".join(tr["excluded_signal_ids"]),
                 size=9, color=MUTED)
        if tr.get("delta"):
            label(doc, t("delta", lang))
            para(doc, tr["delta"])
        doc.add_paragraph()

    # --- Что не подтвердилось -------------------------------------------
    rejected = data.get("rejected") or []
    doc.add_page_break()
    doc.add_heading(t("rejected", lang), level=1)
    if rejected:
        para(doc, t("rejected_intro", lang))
        table(doc, [t("candidate", lang), t("reason", lang), t("signals", lang)],
              [[r.get("name", ""), r.get("reason", ""), ", ".join(r.get("signal_ids", []))] for r in rejected],
              widths=[4.5, 9.0, 2.5])
    else:
        para(doc, t("rejected_none", lang), italic=True, color=MUTED)

    # --- Приложение A ----------------------------------------------------
    doc.add_page_break()
    doc.add_heading(t("appA", lang), level=1)
    table(doc, ["ID", t("date", lang), t("obs", lang), t("ind_source", lang), t("type", lang),
                t("rel", lang), t("tier", lang), t("link", lang), t("review", lang)],
          [[s["id"], s.get("date_event", ""), s.get("headline_fact", ""), s.get("source", ""),
            s.get("source_type", ""), s.get("reliability", ""), T.tier(s), s.get("url") or "—",
            s.get("next_review") or "—"] for s in data["signals"]],
          widths=[1.1, 1.5, 4.2, 2.4, 1.7, 1.0, 1.1, 1.9, 1.3])

    # --- Приложение B ----------------------------------------------------
    doc.add_page_break()
    doc.add_heading(t("appB", lang), level=1)
    for tr in trends:
        doc.add_heading(tr["name"], level=2)
        scores = tr.get("scores") or {}
        evidence = tr.get("score_evidence") or {}
        rows = []
        for key, ru, en, weight in weight_table:
            value = scores.get(key)
            rows.append([L((ru, en)), weight, "N/D" if value is None else value,
                         "—" if value is None else round(weight * float(value) / 5.0, 1),
                         evidence.get(key, "—")])
        table(doc, [t("criterion", lang), t("weight", lang), t("rating", lang),
                    t("contrib", lang), t("fact_ref", lang)],
              rows, widths=[4.6, 1.1, 1.3, 1.3, 7.7])
        c = T.compute(tr, weight_table)
        line = (f"Score {c['score']} / 100 · coverage {c['coverage']}% · "
                f"{L(T.CONFIDENCE[T.confidence_hint(tr, c, index)])}")
        if c["coverage"] < 100:
            line += f" · {t('normalized', lang)} {c['normalized']}"
        para(doc, line, bold=True)

    # --- Приложение C ----------------------------------------------------
    doc.add_page_break()
    doc.add_heading(t("appC", lang), level=1)
    by_type = {}
    for s in data["signals"]:
        by_type.setdefault(s.get("source_type", "—"), set()).add(s.get("source", "—"))
    table(doc, [t("src_type", lang), t("src_used", lang), t("src_count", lang)],
          [[k, ", ".join(sorted(v)), sum(1 for s in data["signals"] if s.get("source_type") == k)]
           for k, v in sorted(by_type.items())], widths=[3.4, 10.4, 2.2])
    tiers = {}
    for s in data["signals"]:
        tiers[T.tier(s)] = tiers.get(T.tier(s), 0) + 1
    para(doc, f"{t('tier_note', lang)}: " + ", ".join(f"{k} — {tiers.get(k, 0)}"
                                                     for k in ("T1", "T2", "T3", "T4")), bold=True)
    para(doc, t("appC_note", lang), italic=True, color=MUTED)

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description="Word-отчёт по трендам из trends.json")
    ap.add_argument("input")
    ap.add_argument("-o", "--output", default="Тренд-отчёт.docx")
    ap.add_argument("--lang", choices=["ru", "en"], default=None)
    args = ap.parse_args()
    data = T.load(args.input)
    path = build(data, args.output, T.lang_of(data, args.lang))
    print(f"Готово: {path}")


if __name__ == "__main__":
    main()
