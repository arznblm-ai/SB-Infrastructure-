"""Приёмка общего сборщика сметы (без сети): смета СИБУР base и два её варианта."""
import json
from pathlib import Path

import pytest

from estimate_build import DEFAULT_TEMPLATE, build

FIXTURE = Path(__file__).parent / "fixtures" / "estimate_sibur_base.json"


def _base():
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def test_sibur_base(tmp_path):
    res = build(_base(), DEFAULT_TEMPLATE, tmp_path / "base.xlsx")
    assert res.subtotal == 1_440_000
    assert res.total == 1_738_800
    assert res.missing_roles == []
    assert res.out_path.exists()


def test_sibur_plus_camera_tracking(tmp_path):
    est = _base()
    est["lines"].append({"role": "camera_tracking", "qty": 1, "days": 2})
    res = build(est, DEFAULT_TEMPLATE, tmp_path / "ct.xlsx")
    assert res.subtotal == 1_480_000
    assert res.total == 1_787_100
    assert res.missing_roles == []


def test_unknown_line_goes_to_missing_without_changing_total(tmp_path):
    est = _base()
    est["unknown"] = [{"title": "Раскадровщик агентства"}]
    res = build(est, DEFAULT_TEMPLATE, tmp_path / "unk.xlsx")
    assert res.missing_roles == ["Раскадровщик агентства"]
    assert res.total == 1_738_800


def test_role_without_rate_in_template_is_reported(tmp_path):
    est = _base()
    est["lines"].append({"role": "storyboard", "qty": 1, "rate": 90000})
    res = build(est, DEFAULT_TEMPLATE, tmp_path / "sb.xlsx")
    assert "storyboard" in res.missing_roles
    assert res.total == 1_738_800  # ставка из JSON не используется


@pytest.mark.parametrize("patch, needle", [
    ({"fee_pct": 1.5}, "fee_pct"),
    ({"tax_pct": -0.1}, "tax_pct"),
    ({"client": ""}, "client"),
])
def test_validation_errors(tmp_path, patch, needle):
    est = _base()
    est.update(patch)
    with pytest.raises(ValueError) as e:
        build(est, DEFAULT_TEMPLATE, tmp_path / "bad.xlsx")
    assert needle in str(e.value)


def test_day_role_requires_days(tmp_path):
    est = _base()
    est["lines"].append({"role": "vfx", "qty": 1})
    with pytest.raises(ValueError) as e:
        build(est, DEFAULT_TEMPLATE, tmp_path / "bad2.xlsx")
    assert "days" in str(e.value)


def test_unknown_lands_in_post_miscellanea_row_29(tmp_path):
    from openpyxl import load_workbook
    est = _base()
    est["unknown"] = [{"title": "Раскадровщик агентства", "qty": 2}]
    res = build(est, DEFAULT_TEMPLATE, tmp_path / "misc.xlsx")
    ws = load_workbook(res.out_path)["POST PRODUCTION"]
    assert ws["C29"].value == "Раскадровщик агентства ?"
    assert ws["E29"].value == 2
    assert ws["G29"].value is None
