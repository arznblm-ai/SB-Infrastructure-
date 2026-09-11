"""Смета «СИБУР мелиорация» (тендер, full-AI ролик 60-75 сек, base) в xlsx-шаблоне Антона.

Шаблон: смета VTB Privilege «Total Dop_2» от 07.09.2026 (Drive id 16ps-kJ8pyl8yH-ccGgv8nnsZWjHAh9aR),
структура MAIN <- PRE PRODUCTION / PRODUCTION / POST PRODUCTION <- EXPENSES.

Что чиню в шаблоне (найдено при проверке формул 10.09.2026):
  - MAIN!E22 налог был захардкожен `=(E21)*4.8%`, ячейка D22 (5%) не использовалась -> формула от D22;
    добавлена строка налогов 15% (правило 14: комиссия ~5% + налоги 15%), TOTAL пересобран.
  - Заголовки секций PRE!B16, PRE!B22, POST!B7, POST!B15, POST!B25 ссылались на пустые MAIN!AI10.. -> текст.
  - POST!K23 суммировал K16:K21 без строки Miscellanea K22 -> K16:K22.
  - PRODUCTION «Shooting» K20..K23 считали итог из Q-ty/Rate напрямую -> единая формула E*(G*H+I*J).
  - Колонка USD (MAIN!F) и курс доллара (E26/F26) удалены - смета РФ-клиенту только в рублях (Антон, 10.09).
  - POST!C16 «Ретушер» ссылался на ставку 2D designer (EXPENSES!E23) -> подпись «2D designer / Моушн-дизайнер».

Использование:
    python3 build_estimate_sibur.py <vtb_template.xlsx> <clean_template_out.xlsx> <estimate_out.xlsx>
затем  python3 ../../Alina/Scripts/inject_values.py <estimate_out.xlsx> <estimate_cached.xlsx>
"""
import sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell.cell import MergedCell

SRC, TPL_OUT, EST_OUT = sys.argv[1], sys.argv[2], sys.argv[3]


def copy_style(src, dst):
    dst.font = copy(src.font); dst.fill = copy(src.fill); dst.border = copy(src.border)
    dst.alignment = copy(src.alignment); dst.number_format = src.number_format


def fix_template(wb):
    m = wb["MAIN"]
    for rng in list(m.merged_cells.ranges):
        if rng.min_row in (23, 24):
            m.unmerge_cells(str(rng))
    # row 24 <- style of old TOTAL row 23; row 23 <- style of row 22
    for col in "BCDEF":
        copy_style(m[f"{col}23"], m[f"{col}24"])
        copy_style(m[f"{col}22"], m[f"{col}23"])
    m["B22"] = "Studio fee / Комиссия студии"; m["D22"] = 0.05; m["E22"] = "=E21*$D$22"; m["F22"] = "=E22/$F$26"
    m["B23"] = "Tax / Налоги"; m["D23"] = 0.15; m["E23"] = "=(E21+E22)*$D$23"; m["F23"] = "=E23/$F$26"
    m["B24"] = "TOTAL (fee & taxes are included) / ВСЕГО (включая комиссию и налоги):"
    m["E24"] = "=SUM(E21:E23)"
    for c in ("D22", "D23"):
        m[c].number_format = "0%"
    m.merge_cells("B24:D24")
    # USD-колонка и курс убраны (решение Антона 10.09): смета РФ-клиенту только в рублях
    from openpyxl.cell.cell import MergedCell as _MC
    for r in range(9, 27):
        if not isinstance(m[f"F{r}"], _MC):
            m[f"F{r}"] = None
    m["E26"] = None
    m.column_dimensions["F"].hidden = True

    pp = wb["PRE PRODUCTION"]
    pp["B16"] = "Preproduction expenses / Предпроизводственные затраты"
    pp["B22"] = "Tests / Тесты"
    post = wb["POST PRODUCTION"]
    post["B7"] = "TC & Edit / Цветокоррекция и монтаж"
    post["B15"] = "CGI & VFX / Графика и визуальные спецэффекты"
    post["B25"] = "Sound / Звук"
    post["K23"] = "=SUM(K16:K22)"
    post["C16"] = "2D designer / Моушн-дизайнер"
    for name in ("PRE PRODUCTION", "PRODUCTION", "POST PRODUCTION"):
        ws = wb[name]
        ws.column_dimensions["C"].width = 38   # названия ролей не режутся примечанием (C:D в шаблоне были слиты)
    prod = wb["PRODUCTION"]
    for r in range(20, 24):
        prod[f"K{r}"] = f"=E{r}*(G{r}*H{r}+I{r}*J{r})"
        prod[f"F{r}"] = "per proj"


def clear_inputs(wb):
    """Очистить количество/дни/часы OT и шапку клиента; ставки и формулы оставить."""
    for name in ("PRE PRODUCTION", "PRODUCTION", "POST PRODUCTION"):
        ws = wb[name]
        for row in ws.iter_rows(min_row=8, max_row=ws.max_row):
            for c in row:
                if isinstance(c, MergedCell):
                    continue
                if c.column in (4, 5, 8, 10) and not (isinstance(c.value, str) and c.value.startswith("=")):
                    c.value = None
    m = wb["MAIN"]
    for c in ("C6", "C7", "C8"):
        m[c] = None


def line(ws, r, qty=None, days=None, note=None):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == r and rng.min_col == 3 and rng.max_col == 4:
            ws.unmerge_cells(str(rng))
    if qty is not None:
        ws.cell(row=r, column=5).value = qty
    if days is not None:
        ws.cell(row=r, column=8).value = days
    if note is not None:
        cell = ws.cell(row=r, column=4)
        cell.value = note
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(name="Roboto", sz=9)
        ws.row_dimensions[r].height = 27 if len(note) > 48 else 15.75


wb = load_workbook(SRC)
fix_template(wb)
clear_inputs(wb)
wb.save(TPL_OUT)
wb = load_workbook(TPL_OUT)  # второе сохранение того же объекта ломает встроенный логотип (закрытый файл картинки)

# ---------------- base: СИБУР мелиорация ----------------
m = wb["MAIN"]

m["C6"] = "СИБУР"
m["C7"] = "Мелиорация - продуктовый AI-ролик 60-75 сек"
m["C8"] = None  # прямой тендер: агентство не названо

pp = wb["PRE PRODUCTION"]
line(pp, 8, 1, 15)
line(pp, 9, 1, 10)
line(pp, 11, 3, 12)

post = wb["POST PRODUCTION"]
line(post, 8, 1, 3)
line(post, 16, 1, 6)
line(post, 26, 1, None)
line(post, 27, 1, None)
line(post, 28, 1, None)

wb.save(EST_OUT)
print("saved", TPL_OUT, EST_OUT)
