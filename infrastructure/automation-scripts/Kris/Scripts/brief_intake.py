#!/usr/bin/env python3
"""Приведение входящего брифа к тексту, который может прочитать модель.

ADR-027: модель Claude в боте Крис не имеет Bash и сети. Всё скачивание,
конвертация и рендер делает код — модуль кладёт результат в
``workdir/inbox/<stamp>/`` и отдаёт готовый блок для промпта.

Модуль НЕ ходит в сеть. Google-файлы забираются через переданный объект
``drive`` (duck typing: метод ``export(file_id, mime) -> Path``).

CLI:
    python3 Scripts/brief_intake.py <file|url|-> --workdir DIR
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

# --- Лимиты -----------------------------------------------------------------

MAX_PAGES = 40
MAX_PNG = 15
IMAGE_PAGE_TEXT_THRESHOLD = 200  # < этого числа символов → страница-картинка
PNG_DPI = 110

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml"}


@dataclass
class IntakeResult:
    kind: str  # pdf | docx | text | gdoc | gslides | gsheet | figma_unsupported | unknown_link | unsupported
    dir: Path
    paths: list[Path] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    summary: str = ""

    def to_json(self) -> dict:
        return {
            "kind": self.kind,
            "dir": str(self.dir),
            "paths": [str(p) for p in self.paths],
            "warnings": list(self.warnings),
            "summary": self.summary,
        }


# --- Вспомогательное --------------------------------------------------------


def _stamp() -> str:
    return datetime.now().strftime("%Y%m%d-%H%M%S")


def _make_dir(workdir: Path) -> Path:
    base = Path(workdir) / "inbox"
    stamp = _stamp()
    target = base / stamp
    n = 1
    while target.exists():
        target = base / f"{stamp}-{n}"
        n += 1
    target.mkdir(parents=True, exist_ok=True)
    return target


def _import_fitz():
    try:
        import fitz  # type: ignore
    except ImportError as exc:  # pragma: no cover - зависит от окружения
        raise RuntimeError(
            "Для разбора PDF нужен пакет pymupdf (import fitz). "
            "Поставь его в окружение бота: pip install pymupdf"
        ) from exc
    return fitz


# --- PDF --------------------------------------------------------------------


def _pdf_to_pages(pdf_path: Path, out_dir: Path) -> IntakeResult:
    fitz = _import_fitz()
    warnings: list[str] = []
    paths: list[Path] = []

    doc = fitz.open(str(pdf_path))
    total = doc.page_count
    limit = min(total, MAX_PAGES)
    if total > MAX_PAGES:
        warnings.append(
            f"В PDF {total} страниц, разобраны первые {MAX_PAGES}. Остальное — попроси у Антона отдельно."
        )

    lines: list[str] = []
    png_made = 0
    png_names: list[str] = []
    for i in range(limit):
        page = doc.load_page(i)
        text = (page.get_text() or "").strip()
        lines.append(f"## Страница {i + 1}")
        if len(text) < IMAGE_PAGE_TEXT_THRESHOLD:
            if png_made < MAX_PNG:
                name = f"page-{i + 1:02d}.png"
                pix = page.get_pixmap(dpi=PNG_DPI)
                png_path = out_dir / name
                pix.save(str(png_path))
                paths.append(png_path)
                png_names.append(name)
                png_made += 1
                if text:
                    lines.append(text)
                lines.append(f"[картинка: {name}]")
            else:
                if text:
                    lines.append(text)
                lines.append("[картинка не отрендерена: превышен лимит PNG]")
        else:
            lines.append(text)
        lines.append("")

    doc.close()

    if png_made >= MAX_PNG:
        warnings.append(
            f"Отрендерено {png_made} PNG (лимит {MAX_PNG}); остальные страницы-картинки без изображения."
        )

    pages_md = out_dir / "pages.md"
    pages_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    paths.insert(0, pages_md)

    png_note = f", картинками отрендерены страницы: {', '.join(png_names)}" if png_names else ""
    summary = (
        f"PDF, страниц: {total} (разобрано {limit}){png_note}. "
        f"Текст — в {pages_md}."
    )
    return IntakeResult(kind="pdf", dir=out_dir, paths=paths, warnings=warnings, summary=summary)


# --- DOCX -------------------------------------------------------------------


def _docx_text_via_python_docx(path: Path) -> str | None:
    try:
        import docx  # type: ignore
    except ImportError:
        return None
    doc = docx.Document(str(path))
    chunks: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            chunks.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))
    return "\n\n".join(chunks)


def _docx_text_via_zip(path: Path) -> str:
    import xml.etree.ElementTree as ET
    import zipfile

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    body = root.find("w:body", ns)
    if body is None:
        return ""

    def para_text(p) -> str:
        return "".join(t.text or "" for t in p.iter(f"{{{ns['w']}}}t")).strip()

    chunks: list[str] = []
    for child in body:
        tag = child.tag.split("}")[-1]
        if tag == "p":
            t = para_text(child)
            if t:
                chunks.append(t)
        elif tag == "tbl":
            for tr in child.findall("w:tr", ns):
                cells = [para_text(tc) for tc in tr.findall("w:tc", ns)]
                if any(cells):
                    chunks.append(" | ".join(cells))
    return "\n\n".join(chunks)


def _docx_to_pages(path: Path, out_dir: Path) -> IntakeResult:
    warnings: list[str] = []
    text = _docx_text_via_python_docx(path)
    if text is None:
        warnings.append("python-docx не установлен, текст извлечён упрощённо (zipfile+xml).")
        text = _docx_text_via_zip(path)
    if not text.strip():
        warnings.append("В документе не найдено текста — возможно, всё содержимое в картинках.")

    pages_md = out_dir / "pages.md"
    pages_md.write_text(f"## Страница 1\n{text}\n", encoding="utf-8")
    summary = f"DOCX, {len(text)} символов текста. Текст — в {pages_md}."
    return IntakeResult(kind="docx", dir=out_dir, paths=[pages_md], warnings=warnings, summary=summary)


# --- Текст ------------------------------------------------------------------


def _write_text_result(text: str, out_dir: Path, kind: str, label: str) -> IntakeResult:
    warnings: list[str] = []
    if not text.strip():
        warnings.append("Пустой текст.")
    pages_md = out_dir / "pages.md"
    pages_md.write_text(f"## Страница 1\n{text.strip()}\n", encoding="utf-8")
    summary = f"{label}, {len(text)} символов. Текст — в {pages_md}."
    return IntakeResult(kind=kind, dir=out_dir, paths=[pages_md], warnings=warnings, summary=summary)


# --- URL --------------------------------------------------------------------

_GOOGLE_PATTERNS = [
    ("gdoc", re.compile(r"docs\.google\.com/document/d/([A-Za-z0-9_-]+)")),
    ("gslides", re.compile(r"docs\.google\.com/presentation/d/([A-Za-z0-9_-]+)")),
    ("gsheet", re.compile(r"docs\.google\.com/spreadsheets/d/([A-Za-z0-9_-]+)")),
]
_FIGMA_RE = re.compile(r"(?:www\.)?figma\.com/(?:file|design|proto|board|slides)/([A-Za-z0-9_-]+)")


def classify_url(url: str) -> tuple[str, str | None]:
    """Определить тип ссылки и id файла. Сеть не трогается."""
    u = (url or "").strip()
    for kind, rx in _GOOGLE_PATTERNS:
        m = rx.search(u)
        if m:
            return kind, m.group(1)
    m = _FIGMA_RE.search(u)
    if m:
        return "figma", m.group(1)
    if "figma.com" in u:
        return "figma", None
    return "other", None


def ingest_url(url: str, workdir: Path, drive=None) -> IntakeResult:
    kind, file_id = classify_url(url)
    out_dir = _make_dir(workdir)
    (out_dir / "source.url").write_text(url.strip() + "\n", encoding="utf-8")

    if kind == "figma":
        return IntakeResult(
            kind="figma_unsupported",
            dir=out_dir,
            paths=[out_dir / "source.url"],
            warnings=["Figma не поддерживается: доступа к API у бота нет."],
            summary="ссылка Figma: попроси у Антона PDF-экспорт",
        )

    if kind == "other":
        return IntakeResult(
            kind="unknown_link",
            dir=out_dir,
            paths=[out_dir / "source.url"],
            warnings=["Ссылка не распознана; скачивание не выполняется (сети у бота нет)."],
            summary=f"Неизвестная ссылка {url.strip()}: содержимое не получено, попроси файл вложением.",
        )

    # Google-файлы
    if drive is None:
        return IntakeResult(
            kind="unknown_link",
            dir=out_dir,
            paths=[out_dir / "source.url"],
            warnings=["Google Drive не настроен — файл не выгружен."],
            summary=(
                f"Google-ссылка ({kind}, id {file_id}): Drive не настроен, "
                "содержимое не получено — попроси файл вложением."
            ),
        )

    if kind == "gdoc":
        exported = Path(drive.export(file_id, "text/plain"))
        target = out_dir / "source.txt"
        if exported.resolve() != target.resolve():
            shutil.copy2(exported, target)
        text = target.read_text(encoding="utf-8", errors="replace")
        res = _write_text_result(text, out_dir, "gdoc", "Google Doc")
        res.paths.insert(0, target)
        return res

    if kind == "gsheet":
        exported = Path(drive.export(file_id, "text/csv"))
        target = out_dir / "source.csv"
        if exported.resolve() != target.resolve():
            shutil.copy2(exported, target)
        text = target.read_text(encoding="utf-8", errors="replace")
        res = _write_text_result(text, out_dir, "gsheet", "Google Sheet (CSV)")
        res.paths.insert(0, target)
        return res

    # gslides → PDF → как PDF
    exported = Path(drive.export(file_id, "application/pdf"))
    target = out_dir / "source.pdf"
    if exported.resolve() != target.resolve():
        shutil.copy2(exported, target)
    res = _pdf_to_pages(target, out_dir)
    res.kind = "gslides"
    res.paths.insert(0, target)
    res.summary = "Google Slides → PDF. " + res.summary
    return res


# --- Публичные входы --------------------------------------------------------


def ingest_text(text: str, workdir: Path) -> IntakeResult:
    out_dir = _make_dir(workdir)
    src = out_dir / "source.txt"
    src.write_text(text, encoding="utf-8")
    res = _write_text_result(text, out_dir, "text", "Текстовый бриф")
    res.paths.insert(0, src)
    return res


def ingest_file(path: Path, workdir: Path, original_name: str | None = None) -> IntakeResult:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {path}")

    name = original_name or path.name
    suffix = Path(name).suffix.lower() or path.suffix.lower()
    out_dir = _make_dir(workdir)
    src = out_dir / f"source{suffix or '.bin'}"
    shutil.copy2(path, src)

    if suffix == ".pdf":
        res = _pdf_to_pages(src, out_dir)
    elif suffix in (".docx", ".dotx"):
        res = _docx_to_pages(src, out_dir)
    elif suffix in TEXT_SUFFIXES:
        text = src.read_text(encoding="utf-8", errors="replace")
        res = _write_text_result(text, out_dir, "text", f"Текстовый файл {name}")
    else:
        return IntakeResult(
            kind="unsupported",
            dir=out_dir,
            paths=[src],
            warnings=[f"Формат {suffix or '(без расширения)'} не поддерживается."],
            summary=(
                f"Файл {name}: формат не поддерживается. "
                "Поддерживаются PDF, DOCX, txt/md. Попроси у Антона PDF-версию."
            ),
        )

    res.paths.insert(0, src)
    return res


# --- Блок для промпта -------------------------------------------------------


def prompt_block(result: IntakeResult) -> str:
    lines = ["=== ВХОДЯЩИЙ БРИФ ==="]
    lines.append(f"Тип: {result.kind}")
    lines.append(f"Папка: {result.dir}")
    lines.append(f"Что это: {result.summary}")
    if result.paths:
        lines.append("Файлы:")
        for p in result.paths:
            lines.append(f"  - {p}")
    if result.warnings:
        lines.append("Предупреждения:")
        for w in result.warnings:
            lines.append(f"  ! {w}")
    if result.kind in ("pdf", "gslides", "docx", "text", "gdoc", "gsheet"):
        lines.append(
            "Инструкция: читай pages.md — это основной текст брифа. "
            "Если в нём встречается пометка [картинка: page-NN.png] — открой соответствующий "
            "PNG из этой же папки, там содержимое страницы-слайда."
        )
    else:
        lines.append(
            "Инструкция: содержимое получить не удалось — скажи об этом Антону "
            "и попроси файл в поддерживаемом формате (PDF/DOCX/текст)."
        )
    lines.append("=== КОНЕЦ БРИФА ===")
    return "\n".join(lines)


# --- CLI --------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Приведение брифа к тексту для модели")
    ap.add_argument("source", help="путь к файлу, URL или '-' для чтения текста из stdin")
    ap.add_argument("--workdir", required=True, help="рабочий каталог (создаст inbox/<stamp>/)")
    ap.add_argument("--original-name", default=None, help="исходное имя файла (для расширения)")
    args = ap.parse_args(argv)

    workdir = Path(args.workdir)
    workdir.mkdir(parents=True, exist_ok=True)

    if args.source == "-":
        result = ingest_text(sys.stdin.read(), workdir)
    elif re.match(r"^https?://", args.source.strip()):
        result = ingest_url(args.source, workdir, drive=None)
    else:
        result = ingest_file(Path(args.source), workdir, original_name=args.original_name)

    print(json.dumps(result.to_json(), ensure_ascii=False, indent=2))
    print(prompt_block(result), file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
