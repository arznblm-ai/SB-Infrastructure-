"""Сравнение сметы-JSON с xlsx, который Антон поправил руками.

Смысл: код собрал смету из `estimate.json`, Антон скачал xlsx с Drive, поправил
количества/дни/добавил позицию. Команда «сверь смету» скачивает файл обратно,
этот модуль считает разницу и отдаёт человекочитаемые строки, которые уходят в
промпт модели - чтобы она записала урок в правила расчёта.

Ставки модель не задаёт и здесь не сравнивает: ставки живут в шаблоне.

Использование:
    from estimate_diff import diff
    lines = diff(json.loads(path.read_text()), Path("правленая.xlsx"))
"""
from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from estimate_build import (
    COL_QTY,
    COL_RATE,
    COL_TITLE,
    COL_DAYS,
    DEFAULT_TEMPLATE,
    MISC_FALLBACK,
    MISC_ROWS,
    ROLE_MAP,
    build,
)

SHEET_SHORT = {
    "PRE PRODUCTION": "PRE",
    "PRODUCTION": "PROD",
    "POST PRODUCTION": "POST",
}

TOTAL_CELL = "E24"


def format_money(value) -> str:
    """1738800.0 -> '1 738 800'. Копейки показываем только если они есть."""
    try:
        num = float(value or 0)
    except (TypeError, ValueError):
        return str(value)
    if abs(num - round(num)) < 0.005:
        body = f"{int(round(num)):,}".replace(",", " ")
    else:
        body = f"{num:,.2f}".replace(",", " ").replace(".", ",")
    return body


def _n(value) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return 0.0
    return float(value)


def _fmt_n(value) -> str:
    num = _n(value)
    return str(int(num)) if abs(num - int(num)) < 1e-9 else f"{num:g}"


def _label(ws, row: int, sheet: str, role: str) -> str:
    title = ws.cell(row=row, column=COL_TITLE).value
    title = str(title).split(" / ")[0].strip() if title else role
    return f"{SHEET_SHORT.get(sheet, sheet)} {title}"


def _expected_book(estimate: dict, template: Path):
    """Смета «как её собрал код»: пересобираем из JSON во временный файл.

    Сравнивать xlsx с xlsx, а не с JSON, - единственный честный путь: сборщик
    сам дописывает H=1 в per-proj строки, и наивное сравнение с JSON давало бы
    ложные расхождения.

    Возвращает (workbook data_only | None, total | None).
    """
    tmp_dir = Path(tempfile.mkdtemp(prefix="kris-diff-"))
    try:
        res = build(estimate, template, tmp_dir / "expected.xlsx")
        return load_workbook(tmp_dir / "expected.xlsx", data_only=True), float(res.total)
    except Exception:  # noqa: BLE001 - диф не должен падать из-за пересборки
        return None, None
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _misc_targets() -> set:
    targets = {(sheet, row) for sheet, row in MISC_ROWS.items()}
    targets.add(MISC_FALLBACK)
    return targets


def diff(estimate_json: dict, xlsx_path, template=None) -> list[str]:
    """Человекочитаемые строки разницы между JSON-сметой и правленым xlsx."""
    if not isinstance(estimate_json, dict):
        raise ValueError("estimate_json: ожидается объект")
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise ValueError(f"файл сметы не найден: {xlsx_path}")
    template = Path(template or DEFAULT_TEMPLATE)

    wb = load_workbook(xlsx_path, data_only=True)
    want_wb, want_total = _expected_book(estimate_json, template)
    changes: list[str] = []

    if want_wb is not None:
        for role, (sheet, row, _kind) in ROLE_MAP.items():
            if sheet not in wb.sheetnames or sheet not in want_wb.sheetnames:
                continue
            ws, want_ws = wb[sheet], want_wb[sheet]
            label = _label(ws, row, sheet, role)
            want_qty = _n(want_ws.cell(row=row, column=COL_QTY).value)
            got_qty = _n(ws.cell(row=row, column=COL_QTY).value)
            want_days = _n(want_ws.cell(row=row, column=COL_DAYS).value)
            got_days = _n(ws.cell(row=row, column=COL_DAYS).value)

            if abs(got_days - want_days) > 1e-9:
                changes.append(
                    f"{label}: было {_fmt_n(want_days)} дн → стало {_fmt_n(got_days)} дн"
                )
            if abs(got_qty - want_qty) > 1e-9:
                changes.append(
                    f"{label}: было {_fmt_n(want_qty)} шт → стало {_fmt_n(got_qty)} шт"
                )

    for sheet, row in sorted(_misc_targets()):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rate = ws.cell(row=row, column=COL_RATE).value
        if rate in (None, ""):
            continue
        title = ws.cell(row=row, column=COL_TITLE).value or "Miscellanea"
        changes.append(f"новая позиция: {str(title).strip()} {format_money(rate)}")

    got_total = _n(wb["MAIN"][TOTAL_CELL].value) if "MAIN" in wb.sheetnames else 0.0
    if want_total is not None and abs(got_total - want_total) > 0.5:
        changes.append(f"итого: {format_money(want_total)} → {format_money(got_total)}")

    return changes
